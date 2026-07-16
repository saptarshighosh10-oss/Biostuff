"""Append paired Roche-34 relative-HIC observations without altering prior data."""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path.home() / "Downloads/antibody_sources/roche_hic_hydrophobicity_scales.xlsx"
DESTINATION = Path.home() / "Downloads/antibody_sources/hydrophobicity_dataset.csv"
FIELDS = ["source", "vh", "vl", "label", "label_type"]
AA = frozenset("ACDEFGHIKLMNPQRSTVWY")


def clean(value: object) -> str | None:
    sequence = str(value or "").strip().upper()
    return sequence if len(sequence) >= 20 and set(sequence) <= AA else None


def main() -> None:
    existing = set()
    with DESTINATION.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            existing.add(row["vh"].strip().upper())

    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    sheet = workbook["SI Table 1"]
    rows = sheet.iter_rows(values_only=True)
    next(rows)  # table title
    headers = next(rows)
    name_i, hic_i, vh_i, vl_i = (headers.index("Name"), headers.index("HIC rel. ret. time (IgG)"),
                                headers.index("H"), headers.index("L"))
    extracted, skipped = [], 0
    for row in rows:
        vh = clean(row[vh_i])
        vl = clean(row[vl_i])
        try:
            hic = float(row[hic_i])
        except (TypeError, ValueError):
            hic = None
        if not vh or not vl or hic is None:
            skipped += 1
            continue
        if vh in existing:
            continue
        existing.add(vh)
        extracted.append({"source": "roche34_2022", "vh": vh, "vl": vl, "label": hic,
                          "label_type": "hic_relative_retention"})
    workbook.close()

    with DESTINATION.open("a", newline="", encoding="utf-8") as handle:
        csv.DictWriter(handle, fieldnames=FIELDS).writerows(extracted)
    values = [row["label"] for row in extracted]
    print(f"roche34_2022: added={len(extracted)}, skipped={skipped}")
    if values:
        print(f"relative HIC range: {min(values):.3f} to {max(values):.3f}; higher = more hydrophobic")


if __name__ == "__main__":
    main()
